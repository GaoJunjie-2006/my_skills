# -*- coding: utf-8 -*-
"""Debug: read exact text from template paragraphs"""
from docx import Document

doc = Document(r"X:\Temp_ws\shiyan4\lab-report-writer\evals\files\测试模板_自动控制原理实验报告.docx")
print("=== ALL PARAGRAPHS (with repr) ===")
for i, para in enumerate(doc.paragraphs):
    t = para.text
    if t.strip():
        print("P%d: %s" % (i, repr(t)))
    else:
        print("P%d: (empty)" % i)

print("\n=== CHECK SPECIFIC TEXT ===")
# Check for the problematic texts
targets = [
    "观察各典型环节",
    "阶跃响应",
    "控制系统由各典型环节组成",
    "示波器",
    "实验一",
    "实验报告到此结束",
]
for target in targets:
    found = False
    for i, para in enumerate(doc.paragraphs):
        if target in para.text:
            print("[FOUND] '%s' in P%d: %s" % (target, i, repr(para.text[:80])))
            found = True
    if not found:
        print("[MISSING] '%s' not found in any paragraph" % target)

print("\n=== CHECK TABLE CELLS ===")
for ti, table in enumerate(doc.tables):
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            t = cell.text.strip()
            if t:
                print("T%d[%d,%d]: %s" % (ti, ri, ci, repr(t)))
