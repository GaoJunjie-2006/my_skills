# -*- coding: utf-8 -*-
"""验证生成的报告内容"""
import os
from docx import Document

path = r"X:\Temp_ws\shiyan4\lab-report-writer-workspace\iteration-1\eval-2\with_skill\outputs\自动化专业实习报告.docx"
doc = Document(path)

print("=" * 60)
print("验证生成文档内容")
print("=" * 60)
print("段落数: %d" % len(doc.paragraphs))
print("表格数: %d" % len(doc.tables))
print()

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        print("[P%d] %s" % (i, text[:100]))
    elif i == 14:  # the deleted paragraph index marker
        print("[P%d] (空 - 第15段已被删除)" % i)

print()
print("--- 表格内容 ---")
for t_idx, table in enumerate(doc.tables):
    print("表格 %d:" % t_idx)
    for r_idx, row in enumerate(table.rows):
        cells = []
        for cell in row.cells:
            cells.append(cell.text.strip())
        print("  行%d: %s" % (r_idx, " | ".join(cells)))
