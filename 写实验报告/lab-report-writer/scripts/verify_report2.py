# -*- coding: utf-8 -*-
"""Verify generated report and write to UTF-8 file"""
import os
from docx import Document

path = r"X:\Temp_ws\shiyan4\lab-report-writer-workspace\iteration-1\eval-2\with_skill\outputs\自动化专业实习报告.docx"
out = r"X:\Temp_ws\shiyan4\lab-report-writer-workspace\iteration-1\eval-2\with_skill\outputs\verified_content.txt"

doc = Document(path)
with open(out, "w", encoding="utf-8") as f:
    f.write("=== 生成报告内容验证 ===\n")
    f.write("段落数: %d, 表格数: %d\n\n" % (len(doc.paragraphs), len(doc.tables)))

    for i, para in enumerate(doc.paragraphs):
        t = para.text.strip()
        if t:
            f.write("P%d: %s\n" % (i, t))
        else:
            f.write("P%d: (空)\n" % i)

    f.write("\n--- 表格 ---\n")
    for ti, table in enumerate(doc.tables):
        f.write("表格%d:\n" % ti)
        for ri, row in enumerate(table.rows):
            cells = [cell.text.strip().replace('\n', '|') for cell in row.cells]
            f.write("  行%d: %s\n" % (ri, " | ".join(cells)))

    f.write("\n--- 检查项 ---\n")
    # 1. 是否包含"请删除"内容
    has_delete_marker = any("请删除" in p.text for p in doc.paragraphs)
    f.write("含[请删除]内容: %s（期望: False）\n" % has_delete_marker)

    # 2. 是否包含占位符
    has_placeholder = any("请在此处" in p.text for p in doc.paragraphs)
    f.write("含占位符[请在此处]: %s（期望: False）\n" % has_placeholder)

    # 3. 信息是否正确
    info_correct = any("张三" in p.text for p in doc.paragraphs)
    f.write("含[张三]: %s（期望: True）\n" % info_correct)
    info_correct2 = any("2023001" in p.text for p in doc.paragraphs)
    f.write("含[2023001]: %s（期望: True）\n" % info_correct2)

    # 4. 表格信息
    has_table_name = False
    has_table_id = False
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "张三" in cell.text:
                    has_table_name = True
                if "2023001" in cell.text:
                    has_table_id = True
    f.write("表格中[张三]: %s（期望: True）\n" % has_table_name)
    f.write("表格中[2023001]: %s（期望: True）\n" % has_table_id)

    # 5. 图片占位符
    img_placeholders = [p.text.strip() for p in doc.paragraphs if "【" in p.text and "图片" in p.text]
    f.write("\n图片占位符:\n")
    for ph in img_placeholders:
        f.write("  %s\n" % ph)

    # 6. 章节标题
    f.write("\n章节标题:\n")
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("一、") or t.startswith("二、") or t.startswith("三、") or \
           t.startswith("四、") or t.startswith("五、") or t.startswith("六、"):
            f.write("  %s\n" % t)

    f.write("\n验证完成\n")

print("Written to verified_content.txt")
