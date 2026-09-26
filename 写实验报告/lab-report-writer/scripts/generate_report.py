# -*- coding: utf-8 -*-
"""
实习报告生成脚本 v5 - 基于模板生成实习报告
使用正确的模板文本进行匹配
"""

import os
from docx import Document

TEMPLATE_PATH = r"X:\Temp_ws\shiyan4\lab-report-writer\evals\files\测试模板_自动控制原理实验报告.docx"
OUTPUT_DIR = r"X:\Temp_ws\shiyan4\lab-report-writer-workspace\iteration-1\eval-2\with_skill\outputs"
LOG_PATH = os.path.join(OUTPUT_DIR, "生成日志.txt")

with open(LOG_PATH, "w", encoding="utf-8") as f:
    f.write("")


def log(msg):
    with open(LOG_PATH, "a", encoding="utf-8") as f:
        f.write(msg + "\n")
    try:
        print(msg)
    except UnicodeEncodeError:
        pass


def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)


def replace_para(para, new_text):
    for run in para.runs:
        run.text = ""
    if para.runs:
        para.runs[0].text = new_text
    else:
        para.add_run(new_text)


def set_cell(cell, text):
    for i, p in enumerate(cell.paragraphs):
        if i == 0:
            p.text = ""
            p.add_run(text)
        else:
            p._element.getparent().remove(p._element)


def delete_table(table):
    table._element.getparent().remove(table._element)


def main():
    log("=" * 60)
    log("实习报告生成脚本 v5")
    log("模板: %s" % TEMPLATE_PATH)
    log("输出: %s" % OUTPUT_DIR)
    log("=" * 60)

    doc = Document(TEMPLATE_PATH)
    log("\n文档: %d段, %d表" % (len(doc.paragraphs), len(doc.tables)))

    # ======================== [1] 删除段落#15（字体说明）=======================
    log("\n[1] 删除段落#15（字体说明）")
    p15 = doc.paragraphs[15]
    log("  原文: '%s'" % p15.text)
    delete_paragraph(p15)
    log("  [OK] 已删除")

    # ======================== [2] 主标题 ========================
    log("\n[2] 修改主标题")
    replace_para(doc.paragraphs[0], "自动化专业实习报告")
    log("  [OK] -> 自动化专业实习报告")

    # ======================== [3] 章节标题 ========================
    log("\n[3] 修改章节标题")
    section_map = {
        "一、实验目的": "一、实习目的",
        "二、实验原理": "二、实习内容",
        "三、实验设备": "三、实习设备与环境",
        "四、实验内容与步骤": "四、实习过程",
        "五、实验数据与结果分析": "五、实习总结与分析",
        "六、实验结论": "六、实习收获与体会",
    }
    for para in doc.paragraphs:
        t = para.text.strip()
        if t in section_map:
            replace_para(para, section_map[t])
            log("  [OK] '%s' -> '%s'" % (t, section_map[t]))

    # ======================== [4] 信息表格 ========================
    log("\n[4] 填写信息表格")
    for table in doc.tables:
        if not any("姓" in c.text and "名" in c.text for row in table.rows for c in row.cells):
            continue
        # Row0: 标题行
        set_cell(table.rows[0].cells[0], "自动化专业实习报告")
        for c in range(1, len(table.rows[0].cells)):
            set_cell(table.rows[0].cells[c], "")
        # Row1: 姓名、学号
        set_cell(table.rows[1].cells[0], "姓    名")
        set_cell(table.rows[1].cells[1], "张三")
        set_cell(table.rows[1].cells[2], "学    号")
        set_cell(table.rows[1].cells[3], "2023001")
        # Row2: 专业、班级
        set_cell(table.rows[2].cells[0], "专    业")
        set_cell(table.rows[2].cells[1], "自动化")
        set_cell(table.rows[2].cells[2], "班    级")
        set_cell(table.rows[2].cells[3], "自动化2301")
        # Row3: 实习信息
        set_cell(table.rows[3].cells[0], "实习单位")
        set_cell(table.rows[3].cells[1], "XX科技有限公司")
        set_cell(table.rows[3].cells[2], "实习周期")
        set_cell(table.rows[3].cells[3], "4周")
        log("  [OK] 信息表格已更新")

    # ======================== [5] 删除实验数据表格 ========================
    log("\n[5] 删除实验数据表格")
    for table in list(doc.tables):
        for cell in table.rows[0].cells:
            if "输入电压" in cell.text:
                delete_table(table)
                log("  [OK] 已删除")
                break

    # ======================== [6] 按段替换内容 ========================
    log("\n[6] 替换段落内容")

    # 使用确认后的模板文本进行匹配
    REPLACEMENTS = {
        # 实验目的条目
        "1. 熟悉典型控制系统的各个环节及其传递函数。": "1. 熟悉西门子S7-1200 PLC控制系统的硬件组成与工作原理",
        "2. 掌握利用运算放大器搭建模拟电路的方法。": "2. 掌握TIA Portal编程软件的使用方法和梯形图程序设计",
        "3. 观察各典型环节的输出响应特性。": "3. 掌握PLC控制系统的现场调试方法与故障排除技巧",
        # 实验原理内容
        "控制系统由各种典型环节组成，常见的典型环节包括比例环节、惯性环节、积分环节、微分环节等。":
            "本次实习主要围绕西门子S7-1200 PLC控制系统展开，涵盖了PLC硬件组态、梯形图编程、HMI画面设计、现场总线配置以及系统联调等环节。实习期间，在工程师指导下参与了实际项目的调试与维护工作。",
        "各环节的传递函数如下：": "实习期间主要涉及的PLC控制系统技术内容如下：",
        # 公式行
        "（1）比例环节：G(s)=K": "（1）西门子S7-1200硬件组态与模块配置",
        "（2）惯性环节：G(s)=K/(Ts+1)": "（2）TIA Portal软件项目创建与程序编写",
        "（3）积分环节：G(s)=K/s": "（3）PLC程序下载、监控与在线调试",
        "（4）微分环节：G(s)=K·s": "（4）HMI触摸屏画面设计与通信配置",
        # 设备条目
        "1. 计算机 x 1台": "1. 西门子S7-1200 PLC（CPU 1214C） x 多套",
        "2. MATLAB/Simulink 仿真软件 x 1套": "2. TIA Portal V17 编程软件 x 1套",
        "3. 数据采集卡 x 1块": "3. 工业触摸屏（KTP700 Basic PN）及各类传感器 x 若干",
        # 实验步骤
        "1. 搭建比例环节的模拟电路": "1. 根据电气原理图完成PLC控制柜的硬件接线",
        "2. 搭建惯性环节的模拟电路": "2. 使用TIA Portal创建项目并进行硬件组态",
        "3. 分别输入阶跃信号，观察并记录输出波形": "3. 编写梯形图程序并下载到PLC进行功能验证",
        # 子标题
        "步骤一：比例环节实验": "（一）PLC硬件系统认识与接线",
        "步骤二：惯性环节实验": "（二）TIA Portal编程与调试",
        # 图片引导句
        "  连接电路图如下所示：": "实习现场PLC控制柜及设备连接情况如下图所示：",
        # 图片占位符
        "  【请插入比例环节电路图】": "【PLC控制系统硬件接线图+图片】",
        "  【请插入惯性环节电路图】": "【TIA Portal编程界面截图+图片】",
        "【请插入实验波形图+图片】": "【PLC程序运行波形图+图片】",
        # 图标题
        "图1 比例环节阶跃响应曲线": "图1 现场设备运行状态图",
        # 数据区域
        "1. 比例环节实验结果": "1. 实习期间完成的主要工作",
        "表1 比例环节实验数据": "表1 实习期间工作内容统计",
        # 结束语
        "— 实验报告结束 —": "— 实习报告结束 —",
    }

    # 占位符内容（需要包含【】的精确匹配）
    PLACEHOLDER_CONTENT = {
        "【请在此处填写具体实验目的】": (
            "本次实习的目的是通过参与XX科技有限公司PLC控制系统的现场调试与维护工作，"
            "将所学的自动化专业理论知识应用于工程实践。具体包括：\n"
            "1. 熟悉工业现场PLC控制系统的硬件组成与工作原理；\n"
            "2. 掌握西门子S7-1200 PLC的编程方法与调试技巧；\n"
            "3. 了解工业自动化系统的现场安装、调试与故障排除流程；\n"
            "4. 培养工程实践能力与职业素养。"
        ),
        "【请在此处补充实验设备信息】": (
            "4. 万用表、示波器等常用检测仪器 x 1套\n"
            "5. 各类传感器（光电、接近、压力传感器）及执行机构 x 若干"
        ),
        "【请在此处总结实验结论，分析实验结果，讨论误差原因】": (
            "通过本次为期4周的实习，我深刻体会到理论知识与工程实践相结合的重要性。"
            "在XX科技有限公司的PLC控制系统调试与维护工作中，我不仅掌握了西门子S7-1200 "
            "PLC的基本编程方法，还学会了使用TIA Portal软件进行项目组态、程序下载和在线调试。\n\n"
            "在实习过程中，我参与了多个现场调试任务，包括控制系统硬件接线检查、PLC程序逻辑验证、"
            "HMI画面功能测试以及系统联调等工作。通过解决实际工程中的故障问题，"
            "我提高了独立分析和解决问题的能力。\n\n"
            "本次实习让我认识到，扎实的专业知识是工程实践的基础，而严谨的工作态度和良好的团队协作能力"
            "则是完成项目的关键保障。这些宝贵的经验将为我未来的职业发展奠定坚实的基础。"
        ),
    }

    # 合并所有替换
    all_replacements = {}
    all_replacements.update(REPLACEMENTS)
    all_replacements.update(PLACEHOLDER_CONTENT)

    # 执行替换
    replaced_count = 0
    for para in doc.paragraphs:
        t = para.text
        if t in all_replacements:
            old_short = t[:60]
            new_text = all_replacements[t]
            replace_para(para, new_text)
            replaced_count += 1
            log("  [%d] '%s' -> '%s'" % (replaced_count, old_short, new_text[:60]))

    log("  共替换 %d 段" % replaced_count)

    # ======================== [7] 处理表格中"实验一"标题 ========================
    log("\n[7] 处理表格中的实验项目标题")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "实验一" in cell.text:
                    set_cell(cell, "自动化专业实习报告")
                    log("  [OK] 替换表格中实验项目标题")

    # ======================== [8] 填入工作内容 ========================
    log("\n[8] 填入实习工作内容")
    for i, para in enumerate(doc.paragraphs):
        if "表1 实习期间工作内容统计" in para.text:
            # 找到后面的空段落
            for j in range(i+1, min(i+5, len(doc.paragraphs))):
                np = doc.paragraphs[j]
                if np.text.strip() == "":
                    replace_para(np, (
                        "实习期间主要完成了以下工作：\n"
                        "（1）参与PLC控制柜的硬件安装与接线，包括电源模块、CPU模块、I/O模块的安装与配线；\n"
                        "（2）使用TIA Portal V17进行项目组态，完成S7-1200的硬件配置和网络设置；\n"
                        "（3）编写梯形图程序，实现电机启停控制、顺序控制和报警功能；\n"
                        "（4）设计HMI触摸屏监控画面，实现设备状态显示和参数设置功能；\n"
                        "（5）参与现场调试，排查并解决了多个系统故障问题。"
                    ))
                    log("  [OK] 填入工作内容到P%d" % j)
                    break
            break

    # ======================== [9] 最终检查 ========================
    log("\n[9] 最终检查")
    remaining = 0
    for i, para in enumerate(doc.paragraphs):
        t = para.text.strip()
        # 检查是否还有模板占位符
        if "请在此处" in t or "请插入" in t or "请删除" in t:
            log("  [注意] P%d 可能还有占位符: '%s'" % (i, t[:60]))
            remaining += 1
        # 检查是否还有"实验"关键词但不应保留的
        if "实验" in t and "实习" not in t and "实验报告" not in t and "实验设备" not in t:
            # 可能没问题（如"实验设备"是设备列表的一部分）
            pass
    if remaining == 0:
        log("  [OK] 无残留占位符")

    # ======================== 保存 ========================
    log("\n" + "=" * 60)
    output_path = os.path.join(OUTPUT_DIR, "自动化专业实习报告.docx")
    doc.save(output_path)
    log("文件: %s" % output_path)
    log("=" * 60)
    if os.path.exists(output_path):
        log("大小: %.1f KB" % (os.path.getsize(output_path) / 1024))
        log("生成成功！")
    else:
        log("错误：文件未成功保存！")

    return output_path


if __name__ == "__main__":
    main()
